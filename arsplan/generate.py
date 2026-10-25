"""Genererer årsplanen: JSON, masterprompt (Markdown), XLSX og DOCX.

Kjør:  python3 arsplan/generate.py
Krever: python-docx, openpyxl
"""

import json
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from plan_data import (
    CONFIDENTIAL_WEEKS, DAY_365, PHASES, ROLE_MANDATES, ROLES, SKILLS,
    START_DATE, TYPES, WEEKS,
)

OUT = Path(__file__).parent
WEEKDAYS = ["man", "tir", "ons", "tor", "fre", "lør", "søn"]
MONTHS = ["jan", "feb", "mar", "apr", "mai", "jun", "jul", "aug", "sep", "okt", "nov", "des"]
STATUSES = ["Ikke startet", "Pågår", "Venter godkjenning", "Blokkert", "Ferdig", "Droppet"]
MAX_REPAIRS = 2


def tid(n):
    return f"ÅR-{n:03d}"


def phase_for(week):
    for lo, hi, name, _ in PHASES:
        if lo <= week <= hi:
            return name
    return "Dag 365"


def fmt_date(d):
    return f"{WEEKDAYS[d.weekday()]} {d.day}. {MONTHS[d.month - 1]} {d.year}"


def parse(raw):
    typ, title = raw.split(":", 1)
    approval = title.endswith(" [G]")
    return typ, title.removesuffix(" [G]").strip(), approval


def build_tasks():
    start = date.fromisoformat(START_DATE)
    raw = [(w + 1, d + 1, theme, item)
           for w, (theme, items) in enumerate(WEEKS)
           for d, item in enumerate(items)]
    raw.append((53, 1, "Årsavslutning", DAY_365))
    assert len(raw) == 365, len(raw)

    tasks = []
    for n, (week, wday, theme, item) in enumerate(raw, start=1):
        typ, title, approval = parse(item)
        spec = TYPES[typ]
        confidential = spec.get("konfidensiell", False) or week in CONFIDENTIAL_WEEKS
        quarterly = theme.startswith("Kvartalsrevisjon") or n == 365
        weekly_review = typ == "A" and title.startswith("Ukesrevisjon")

        if n == 365:
            deps = [tid(i) for i in range(1, 365)]
        elif weekly_review:
            deps = [tid(i) for i in range(n - wday + 1, n)]
        elif n > 1:
            deps = [tid(n - 1)]
        else:
            deps = []

        if typ in "BOXP" or approval or quarterly or n == 1:
            prio = "P1"
        elif weekly_review or typ == "K":
            prio = "P3"
        else:
            prio = "P2"

        subagents = [{"rolle": r, "navn": ROLES[r][0], "oppdrag": ROLE_MANDATES[r].format(t=title)}
                     for r in spec["roller"]]
        control = {"rolle": "KTR", "navn": ROLES["KTR"][0], "oppdrag": ROLE_MANDATES["KTR"].format(t=title)}
        steps = [s.format(t=title) for s in spec["deloppgaver"]]
        criteria = list(spec["kriterier"])
        if approval:
            criteria.append("Menneskelig godkjenning er registrert i state.json før utførelse")
        if confidential:
            criteria.append("Ingen del av leveransen er delt utenfor lukket mappe")

        cost = spec["kost"] * (2 if quarterly else 1)
        inputs = ["Gjeldende state.json og oppgaveregister"]
        if deps and n != 365:
            inputs.append("Leveranser fra " + (", ".join(deps) if len(deps) <= 6 else f"{deps[0]}–{deps[-1]}"))
        elif n == 365:
            inputs.append("Leveranser fra ÅR-001–364")
        inputs.append("Relevante notater i kunnskapslageret (når det finnes)")

        rules = {
            "kostnadstak_usd": cost,
            "konfidensiell": confidential,
            "krever_godkjenning": approval,
            "maks_reparasjonsrunder": MAX_REPAIRS,
            "regler": [
                f"Stopp og marker «Blokkert» hvis forbruket passerer {cost} USD.",
                "Konfidensielt: kun lukket mappe, ingen eksterne verktøy som lagrer innhold, ingen publisering."
                if confidential else "Ingen publisering, kjøp eller ekstern kontakt uten godkjenning.",
                "Pause ved godkjenningsport; lagre tilstand og vent." if approval
                else "Ingen godkjenningsport; kjør til ferdig eller blokkert.",
            ],
        }
        next_id = tid(n + 1) if n < 365 else None
        state = {
            "ved_ferdig": f"Sett {tid(n)}=Ferdig, logg artefaktstier, neste = {next_id or 'ingen (året er avsluttet)'}.",
            "ved_blokkert": f"Etter {MAX_REPAIRS} reparasjonsrunder: sett {tid(n)}=Blokkert med årsak, "
                            "gå til neste oppgave uten blokkerte avhengigheter.",
        }

        d = start + timedelta(days=n - 1)
        task = {
            "id": tid(n), "nr": n, "dag": n, "uke": min(week, 52), "ukedag": wday,
            "dato": d.isoformat(), "dato_tekst": fmt_date(d),
            "fase": phase_for(week), "ukestema": theme,
            "type": typ, "typenavn": spec["navn"], "tittel": title,
            "prioritet": prio,
            "arbeidsordre": f"{title}. Lever: {spec['leveranse'].lower()}.",
            "subagenter": subagents, "kontrollor": control,
            "deloppgaver": steps, "skills": spec["skills"],
            "input": inputs, "avhengigheter": deps,
            "leveranse": spec["leveranse"], "ferdigkriterier": criteria,
            "regler": rules, "tilstand": state,
            "status": "Ikke startet",
        }
        task["agentarbeidsordre"] = work_order(task)
        tasks.append(task)
    return tasks


def work_order(t):
    r = t["regler"]
    dep = t["avhengigheter"]
    dep_txt = "ingen" if not dep else (", ".join(dep) if len(dep) <= 6 else f"{dep[0]}–{dep[-1]} ({len(dep)} oppgaver)")
    lines = [
        f"# Arbeidsordre {t['id']} — {t['tittel']}",
        f"Dag {t['dag']} · uke {t['uke']} · {t['dato_tekst']} · {t['fase']} · {t['prioritet']}",
        "",
        "Du er Orkestrator. Les masterprompten og state.json før du gjør noe.",
        "",
        f"## Mål\n{t['arbeidsordre']}",
        "",
        "## Avhengigheter\n" + ("Ingen. Kan starte umiddelbart." if not dep
                                 else f"{dep_txt}. Hvis noen av dem ikke er «Ferdig», stopp og rapporter."),
        "",
        "## Input",
        *[f"- {i}" for i in t["input"]],
        "",
        "## Subagenter (separate kjøringer, egne resultater)",
        *[f"{i}. {s['navn']} ({s['rolle']}): {s['oppdrag']}" for i, s in enumerate(t["subagenter"], 1)],
        f"Deretter uavhengig {t['kontrollor']['navn']} (KTR): {t['kontrollor']['oppdrag']} "
        "Kontrolløren ser bare leveransen og ferdigkriteriene, ikke arbeidsprosessen.",
        "",
        "## Deloppgaver i rekkefølge",
        *[f"{i}. {s}" for i, s in enumerate(t["deloppgaver"], 1)],
        "",
        "## Skills",
        *[f"- {s} {SKILLS[s][0]}" for s in t["skills"]],
        "",
        f"## Leveranse\n{t['leveranse']}",
        "",
        "## Ferdigkriterier",
        *[f"- {c}" for c in t["ferdigkriterier"]],
        "",
        "## Regler",
        f"- Kostnadstak: {r['kostnadstak_usd']} USD",
        f"- Konfidensiell: {'ja' if r['konfidensiell'] else 'nei'}",
        f"- Krever godkjenning: {'ja' if r['krever_godkjenning'] else 'nei'}",
        f"- Maks reparasjonsrunder: {r['maks_reparasjonsrunder']}",
        *[f"- {x}" for x in r["regler"]],
        "",
        "## Tilstand og neste handling",
        f"- {t['tilstand']['ved_ferdig']}",
        f"- {t['tilstand']['ved_blokkert']}",
    ]
    return "\n".join(lines)


MASTERPROMPT = """# Masterprompt — årsmaskinen

Du er **Orkestrator** for et produksjonsår med 365 oppdrag (ÅR-001 til ÅR-365).
Du produserer ikke fagarbeid selv. Du delegerer, syntetiserer, kontrollerer og lagrer tilstand.

## 1. Før hver kjøring
1. Les `state.json`. Den er sannheten, ikke hukommelsen din.
2. Finn første oppgave med status «Ikke startet» eller «Pågår» der alle avhengigheter er «Ferdig».
3. Les oppgavens arbeidsordre (`agentarbeidsordre` i JSON). Følg den bokstavelig.

## 2. Arbeidsløkken
Oppgave → avgrensede subagenter → produksjon → syntese → uavhengig kontroll → eventuell reparasjon → lagret tilstand → neste klare oppgave.

- Hver fagrolle får **egen** arbeidsordre og **egen** kjøring. Ikke simuler roller i én lang tekst.
- Gi hver subagent bare det den trenger: mål, input, format, ferdigkriterier.
- Syntetiser resultatene til én leveranse.
- Send leveransen og ferdigkriteriene — ikke prosessen — til Kvalitetskontrolløren.
- Returnert? Reparer og kontroller på nytt. **Maks {repairs} reparasjonsrunder**, deretter «Blokkert» med årsak.

## 3. Godkjenningsporter
Oppgaver med `krever_godkjenning: true` skal:
1. Forberede alt til og med siste steg før det irreversible (publisering, kjøp, ekstern kontakt, innsending).
2. Sette status «Venter godkjenning» og lagre tilstand.
3. Stoppe. Fortsett først når godkjenning er registrert i `state.json`.

## 4. Konfidensialitet (oppfinnelsessporet)
Oppgaver med `konfidensiell: true`:
- Bare lukket mappe. Ingen publisering, ingen deling, ingen verktøy som lagrer eller trener på innholdet.
- Offentliggjøring før innlevering kan bli kjent teknikk og ødelegge nyhet (EPC art. 54). Vurderinger er indikative, ikke juridisk råd.

## 5. Kostnader
- Hver oppgave har `kostnadstak_usd`. Logg forbruk per subagent.
- Passerer forbruket taket: stopp, marker «Blokkert», skriv hvorfor.

## 6. Tilstand
Etter hver oppgave skriver du til `state.json`:
```json
{{"id": "ÅR-NNN", "status": "Ferdig|Blokkert|Venter godkjenning|Droppet",
  "artefakter": ["sti/til/fil"], "kostnad_usd": 0.0,
  "reparasjonsrunder": 0, "notat": "", "neste": "ÅR-NNN"}}
```
En ny sesjon skal kunne fortsette fra `state.json` alene, uten å gjenta ferdig arbeid.

## 7. Prioriteringsregel
Videreutvikle det som virker, fremfor å starte nytt. Det er bedre å ende året med tre ting som lever enn 365 halvferdige ideer.
Oppgaver kan droppes eller erstattes ved kvartalsrevisjon; registrer begrunnelsen.

## 8. Roller
{roles}

## 9. Skills
{skills}

## 10. Hva du aldri gjør
- Publiserer, kjøper, kontakter eller sender inn uten registrert godkjenning.
- Hevder at noe er ferdig uten at kontrolløren har godkjent det.
- Deler konfidensielt materiale.
- Fortsetter etter passert kostnadstak.
- Starter en oppgave med uferdige avhengigheter.

Start: ÅR-001 — kartlegg materialet som faktisk er tilgjengelig.
"""


def masterprompt():
    roles = "\n".join(f"- **{k} {v[0]}** — {v[1]}" for k, v in ROLES.items())
    skills = "\n".join(f"- **{k} {v[0]}** — {v[1]}" for k, v in SKILLS.items())
    return MASTERPROMPT.format(repairs=MAX_REPAIRS, roles=roles, skills=skills)


# ---------------------------------------------------------------- XLSX

HEAD_FILL = PatternFill("solid", fgColor="2F4F4F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

COLUMNS = [
    ("Nr", 6, lambda t: t["nr"]),
    ("ID", 9, lambda t: t["id"]),
    ("Dato", 11, lambda t: t["dato"]),
    ("Uke", 5, lambda t: t["uke"]),
    ("Fase", 22, lambda t: t["fase"]),
    ("Ukestema", 28, lambda t: t["ukestema"]),
    ("Type", 16, lambda t: t["typenavn"]),
    ("Tittel", 50, lambda t: t["tittel"]),
    ("Prioritet", 9, lambda t: t["prioritet"]),
    ("Status", 18, lambda t: t["status"]),
    ("Avhengigheter", 18, lambda t: ", ".join(t["avhengigheter"]) if len(t["avhengigheter"]) <= 6
        else f"{t['avhengigheter'][0]}–{t['avhengigheter'][-1]}"),
    ("Subagenter", 34, lambda t: ", ".join(s["navn"] for s in t["subagenter"]) + " → Kvalitetskontrollør"),
    ("Skills", 22, lambda t: ", ".join(t["skills"])),
    ("Deloppgaver", 50, lambda t: "\n".join(f"{i}. {s}" for i, s in enumerate(t["deloppgaver"], 1))),
    ("Leveranse", 40, lambda t: t["leveranse"]),
    ("Ferdigkriterier", 45, lambda t: "\n".join(f"• {c}" for c in t["ferdigkriterier"])),
    ("Kostnadstak USD", 10, lambda t: t["regler"]["kostnadstak_usd"]),
    ("Faktisk kostnad USD", 10, lambda t: None),
    ("Konfidensiell", 10, lambda t: "Ja" if t["regler"]["konfidensiell"] else "Nei"),
    ("Godkjenning", 10, lambda t: "Ja" if t["regler"]["krever_godkjenning"] else "Nei"),
    ("Reparasjonsrunder", 10, lambda t: 0),
    ("Artefakter", 30, lambda t: None),
    ("Notat", 30, lambda t: None),
]


def header(ws, names, widths):
    ws.append(names)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_xlsx(tasks, path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Oppgaver"
    header(ws, [c[0] for c in COLUMNS], [c[1] for c in COLUMNS])
    for t in tasks:
        ws.append([f(t) for _, _, f in COLUMNS])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    last = len(tasks) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    col = {name: get_column_letter(i) for i, (name, _, _) in enumerate(COLUMNS, 1)}
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{col['Status']}2:{col['Status']}{last}")

    # Dagsvisning: velg oppgavenummer i B2.
    dag = wb.create_sheet("Dag")
    dag.column_dimensions["A"].width = 22
    dag.column_dimensions["B"].width = 90
    dag["A1"], dag["A1"].font = "Dagsvisning", Font(bold=True, size=14)
    dag["A2"], dag["B2"] = "Oppgavenummer (1–365):", 1
    dag["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    dag["B2"].font = Font(bold=True, size=12)
    nv = DataValidation(type="whole", operator="between", formula1="1", formula2="365")
    dag.add_data_validation(nv)
    nv.add("B2")
    fields = ["ID", "Dato", "Uke", "Fase", "Ukestema", "Type", "Tittel", "Prioritet", "Status",
              "Avhengigheter", "Subagenter", "Skills", "Deloppgaver", "Leveranse", "Ferdigkriterier",
              "Kostnadstak USD", "Konfidensiell", "Godkjenning", "Notat"]
    for i, f in enumerate(fields, start=4):
        dag.cell(row=i, column=1, value=f).font = Font(bold=True)
        c = dag.cell(row=i, column=2,
                     value=f'=IFERROR(INDEX(Oppgaver!${col[f]}$2:${col[f]}${last},'
                           f'MATCH($B$2,Oppgaver!$A$2:$A${last},0)),"")')
        c.alignment = WRAP
    dag.cell(row=len(fields) + 5, column=1,
             value="Full arbeidsordre for hver oppgave ligger i aarsplan-365.json (feltet agentarbeidsordre).").font = Font(italic=True)

    # Oversikt
    ov = wb.create_sheet("Oversikt")
    header(ov, ["Status", "Antall"], [22, 10])
    S = f"Oppgaver!${col['Status']}$2:${col['Status']}${last}"
    for s in STATUSES:
        ov.append([s, f'=COUNTIF({S},"{s}")'])
    ov.append(["Totalt", f"=COUNTA(Oppgaver!$A$2:$A${last})"])
    ov.append([])
    ov.append(["Kostnadstak sum USD", f"=SUM(Oppgaver!${col['Kostnadstak USD']}$2:${col['Kostnadstak USD']}${last})"])
    ov.append(["Faktisk kostnad sum USD", f"=SUM(Oppgaver!${col['Faktisk kostnad USD']}$2:${col['Faktisk kostnad USD']}${last})"])
    ov.append(["Andel ferdig", f'=IFERROR(B6/B8,0)'])
    ov["B12"].number_format = "0.0%"
    ov.append([])
    ov.append(["Fase", "Antall"])
    for _, _, name, _ in PHASES:
        ov.append([name, f'=COUNTIF(Oppgaver!${col["Fase"]}$2:${col["Fase"]}${last},"{name}")'])
    ov.append(["Dag 365", f'=COUNTIF(Oppgaver!${col["Fase"]}$2:${col["Fase"]}${last},"Dag 365")'])

    ro = wb.create_sheet("Roller")
    header(ro, ["Kode", "Rolle", "Mandat"], [8, 26, 90])
    for k, (n, d) in ROLES.items():
        ro.append([k, n, d])

    sk = wb.create_sheet("Skills")
    header(sk, ["ID", "Skill", "Beskrivelse"], [8, 28, 90])
    for k, (n, d) in SKILLS.items():
        sk.append([k, n, d])

    wb.save(path)


# ---------------------------------------------------------------- DOCX

def write_docx(tasks, mp, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)

    t = doc.add_heading("Årsplan: 365 agentoppdrag", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"Startdato (forslag): {fmt_date(date.fromisoformat(START_DATE))}. "
                          "Datoene er flyttbare. Ingen kjøring, bot, publisering eller kjøp aktiveres av dette dokumentet.")
    p.runs[0].italic = True

    doc.add_heading("Slik brukes planen", 1)
    for s in [
        "Start med masterprompten (kapittel 4) og ÅR-001.",
        "Hent full arbeidsordre for en oppgave fra aarsplan-365.json (feltet agentarbeidsordre) eller fra kapittel 5.",
        "Oppdater status, kostnad og artefakter i regnearket aarsplan-365.xlsx (arket «Oppgaver»). Arket «Dag» viser én oppgave.",
        f"Maks {MAX_REPAIRS} reparasjonsrunder før en oppgave markeres «Blokkert».",
        "Videreutvikle det som virker. Dropp og erstatt oppgaver ved kvartalsrevisjon, med begrunnelse.",
    ]:
        doc.add_paragraph(s, style="List Number")

    doc.add_heading("1. Faser", 1)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Light Grid Accent 1"
    tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text = "Periode", "Hovedarbeid"
    for lo, hi, name, desc in PHASES:
        r = tbl.add_row().cells
        r[0].text, r[1].text = f"Uke {lo}–{hi} · {name}", desc
    r = tbl.add_row().cells
    r[0].text, r[1].text = "Dag 365", "Avstemme alle oppgavene mot faktiske leveranser; årsrapport, eksport og forslag til neste mandat."

    doc.add_heading("2. Roller (14)", 1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Light Grid Accent 1"
    for c, v in zip(tbl.rows[0].cells, ["Kode", "Rolle", "Mandat"]):
        c.text = v
    for k, (n, d) in ROLES.items():
        r = tbl.add_row().cells
        r[0].text, r[1].text, r[2].text = k, n, d

    doc.add_heading("3. Skill-bibliotek (30)", 1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Light Grid Accent 1"
    for c, v in zip(tbl.rows[0].cells, ["ID", "Skill", "Beskrivelse"]):
        c.text = v
    for k, (n, d) in SKILLS.items():
        r = tbl.add_row().cells
        r[0].text, r[1].text, r[2].text = k, n, d

    doc.add_heading("4. Masterprompt", 1)
    for line in mp.splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], 2)
        elif line.strip():
            para = doc.add_paragraph(line.replace("**", ""))
            if line.startswith(("{", "  ", "```")):
                para.runs[0].font.name = "Consolas"

    doc.add_heading("5. De 365 oppdragene", 1)
    cur_week = None
    grey = RGBColor(0x55, 0x55, 0x55)
    for t in tasks:
        if (t["uke"], t["ukestema"]) != cur_week:
            cur_week = (t["uke"], t["ukestema"])
            label = "Dag 365 · Årsavslutning" if t["nr"] == 365 else f"Uke {t['uke']} · {t['ukestema']}"
            doc.add_heading(label, 2)
        doc.add_heading(f"{t['id']} — {t['tittel']}", 3)
        meta = doc.add_paragraph(
            f"Dag {t['dag']} · {t['dato_tekst']} · {t['typenavn']} · {t['prioritet']}"
            + (" · KONFIDENSIELL" if t["regler"]["konfidensiell"] else "")
            + (" · GODKJENNINGSPORT" if t["regler"]["krever_godkjenning"] else ""))
        meta.runs[0].font.color.rgb = grey
        meta.runs[0].font.size = Pt(9)

        def field(label, text):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            para.add_run(label + ": ").bold = True
            para.add_run(text)

        field("Arbeidsordre", t["arbeidsordre"])
        field("Subagenter", "; ".join(f"{s['navn']} — {s['oppdrag']}" for s in t["subagenter"]))
        field("Kontroll", t["kontrollor"]["oppdrag"])
        field("Deloppgaver", " ".join(f"({i}) {s}" for i, s in enumerate(t["deloppgaver"], 1)))
        field("Skills", ", ".join(f"{s} {SKILLS[s][0]}" for s in t["skills"]))
        dep = t["avhengigheter"]
        field("Avhengigheter", "ingen" if not dep else (", ".join(dep) if len(dep) <= 6 else f"{dep[0]}–{dep[-1]}"))
        field("Leveranse", t["leveranse"])
        field("Ferdig når", "; ".join(t["ferdigkriterier"]))
        field("Regler", f"tak {t['regler']['kostnadstak_usd']} USD; " + " ".join(t["regler"]["regler"]))
        field("Neste", t["tilstand"]["ved_ferdig"])

    doc.save(path)


def main():
    tasks = build_tasks()
    mp = masterprompt()
    payload = {
        "versjon": 1, "startdato": START_DATE, "antall_oppgaver": len(tasks),
        "maks_reparasjonsrunder": MAX_REPAIRS, "statuser": STATUSES,
        "faser": [{"uker": f"{lo}-{hi}", "navn": n, "beskrivelse": d} for lo, hi, n, d in PHASES],
        "roller": {k: {"navn": n, "mandat": d} for k, (n, d) in ROLES.items()},
        "skills": {k: {"navn": n, "beskrivelse": d} for k, (n, d) in SKILLS.items()},
        "oppgaver": tasks,
    }
    (OUT / "aarsplan-365.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "masterprompt.md").write_text(mp, encoding="utf-8")
    write_xlsx(tasks, OUT / "aarsplan-365.xlsx")
    write_docx(tasks, mp, OUT / "aarsplan-365.docx")
    print(f"{len(tasks)} oppgaver skrevet til {OUT}")


if __name__ == "__main__":
    main()
